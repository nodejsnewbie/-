# -*- coding: utf-8 -*-
"""校验生成的工作簿：结构、题目数、下拉、以及重算后的公式结果。"""
from pathlib import Path

from openpyxl import load_workbook

docs = Path(__file__).resolve().parent.parent / "docs"
src = docs / "华农智服-需求澄清与确认问卷.xlsx"
recalc = docs / "_recalc.xlsx"

print("=" * 62)
print("A. 原文件结构")
wb = load_workbook(src, data_only=False)
print("sheets:", wb.sheetnames)

total_dv = 0
total_q = 0
for name in wb.sheetnames:
    if not name.startswith("L"):
        continue
    ws = wb[name]
    # 统计下拉区间覆盖的单元格数
    dv_cells = 0
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            dv_cells += (rng.max_row - rng.min_row + 1) * (rng.max_col - rng.min_col + 1)
    # 统计真实题目行：A 列以 L 开头
    q = [c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("L")]
    total_dv += dv_cells
    total_q += len(q)
    print(f"  {name:22s} 题目 {len(q):2d} 行  下拉覆盖 {dv_cells:2d} 格  "
          f"行号 {q[0]}..{q[-1]}")
print(f"  合计题目 {total_q}，下拉覆盖 {total_dv} 格（应相等）")

print("=" * 62)
print("B. 抽样内容")
ws = wb["L1-需求本质"]
for r in (5, 6):
    print(f"  [{ws.cell(r,1).value}] {ws.cell(r,2).value}")
    print(f"      我们的理解: {ws.cell(r,3).value[:40]}...")
    print(f"      重要度: {ws.cell(r,7).value}")

print("=" * 62)
print("C. 重算后的进度表（答案全空，应全为 0）")
wr = load_workbook(recalc, data_only=True)["填写进度"]
for row in wr.iter_rows(min_row=1, max_row=7, max_col=5, values_only=True):
    print("  ", row)

print("=" * 62)
print("D. 重算后公式是否保留")
wf = load_workbook(recalc, data_only=False)["填写进度"]
print("  C2:", wf["C2"].value)
print("  E2:", wf["E2"].value)
print("  C6:", wf["C6"].value)
