# -*- coding: utf-8 -*-
"""
生成《华农智服 · 需求澄清与确认问卷》Excel 工作簿。

设计要点：
- 按「需求深度」分层（L1 需求本质 → L2 业务与流程 → L3 规则与算法 → L4 数据与集成），
  上一层确认后再填下一层，避免客户一次面对几十个细节问题。
- 每题都给出「我们的理解（供参考）」，客户只需判「是否一致」并填写修正，确认成本最低。
- 「是否一致」列做下拉，方便统计与回收。
- 输出到 docs/ 目录。
"""
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BRAND = "004425"
BRAND_TINT = "E8F0EA"
FILL_IN = "FFF9E3"          # 待客户填写的列
GREY_TINT = "F4F6F4"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [
    ("序号", 9),
    ("问题", 38),
    ("我们的理解（供参考，可直接确认或改写）", 50),
    ("是否一致", 15),
    ("请填写你的答案 / 修正", 40),
    ("补充说明（可选）", 24),
    ("重要度", 9),
]

PRIORITY_COLOR = {"必答": "C00000", "建议答": "BF6000", "可选": "808080"}

# ---------------------------------------------------------------- 题目数据
SHEETS = [
    {
        "name": "L1-需求本质",
        "title": "第一层：需求本质",
        "note": "★ 请先填写本表。这一层只谈「为什么做、给谁做、现在怎么做、做成什么样」，"
                "不涉及任何明细规则。本表确认后，我们再发下一层。共 7 题。",
        "groups": [
            ("一、为什么做、给谁做", [
                ("L1-01", "这个平台，如果只能先解决一件事，您最想解决的是哪一件？",
                 "会议纪要里提到「扩大宣传辐射范围」和「提高服务响应影响力」两个方向，"
                 "但没说哪个更急、是不是同一件事。", "必答"),
                ("L1-02", "平台上线后，主要哪些岗位的人会每天使用？由谁最终拍板决定需求？",
                 "我们目前按「只有两个岗位：企业运营管理员、区域督导」来设想，"
                 "并且认为不单独设资质审核员和财务角色。", "必答"),
                ("L1-03", "现在没有这个系统的时候，这些事是怎么做的？"
                          "（农户怎么找到技师、服务完谁记一笔、钱怎么结）",
                 "我们不了解现状流程。目前系统里的流程大部分是我们推测出来的，"
                 "这一题最能纠正我们的理解。", "必答"),
            ]),
            ("二、做成什么样", [
                ("L1-04", "平台要包含哪几类服务？",
                 "我们目前按五类设想：① 上门植保诊断 ② 农资上门配送与维保 "
                 "③ 飞防（无人机）作业 ④ 农机巡检维保 ⑤ 土壤检测。", "必答"),
                ("L1-05", "第一个版本上线时，您希望先把哪一段跑通？",
                 "不确定。两种可能：先管好技师与资质，或先跑通农户下单与派单。", "必答"),
                ("L1-06", "一年以后，平台做到什么程度您认为就算成功了？",
                 "纪要只有方向（辐射范围、服务响应影响力），没有可衡量的标准。", "必答"),
                ("L1-07", "有没有硬性约束？（上线时间、监管检查要求、预算上限）",
                 "只知道「地图定位因年费约 5 万暂不做」，其它约束不清楚。", "必答"),
            ]),
        ],
    },
    {
        "name": "L2-业务范围与流程",
        "title": "第二层：业务范围与主流程",
        "note": "建议在 L1 确认后再填本表。本层确定「每类服务怎么跑、三个端各干什么、谁和谁配合」。"
                "共 14 题。",
        "groups": [
            ("一、服务与下单", [
                ("L2-01", "每一类服务，从接到需求到完成的完整流程是什么？",
                 "不了解。目前系统里的流程是推测的。", "必答"),
                ("L2-02", "农户通过什么方式下单 / 报单？需要提供哪些信息？",
                 "我们按「作物、面积、病虫害症状、照片、位置」设想。", "必答"),
                ("L2-03", "服务怎么定价？（按亩 / 按次 / 按工时 / 其它）",
                 "不清楚。定价方式会直接影响结算口径。", "必答"),
                ("L2-04", "一个技师一天大概能接几单？有没有上限？",
                 "原型里写「每天最多 4 单」，那个数字是猜的。", "建议答"),
            ]),
            ("二、派单与履约", [
                ("L2-05", "派单是人工挑人，还是系统推荐后人工确认？",
                 "我们按「系统按网格内候选排序、人工确认后派单」设想。", "建议答"),
                ("L2-06", "技师可以拒单吗？", "不清楚。", "可选"),
                ("L2-07", "服务「完成」以什么为准？（农户签字 / 拍照 / 系统点完成）",
                 "我们倾向「农户签字验收」。", "必答"),
                ("L2-08", "农户签字验收这个功能，本期就要做吗？",
                 "不确定。它是分红结算的触发点——不做的话本期就没法自动结算。", "必答"),
                ("L2-09", "电子处方（开方）是真实存在的环节吗？必须先开方才能算完成吗？",
                 "我们猜测存在，且是结算的前置条件（不开方不算完成）。", "必答"),
                ("L2-10", "服务过程中，技师需要在系统里上传什么？"
                          "（到场打卡、现场照片、用药量与配方等）",
                 "我们按「到场打卡 + 现场照片」设想。", "建议答"),
                ("L2-11", "什么情况算异常工单？出现了怎么处理？",
                 "原型里写「接单超时预警后紧急改派」，那是猜的。", "建议答"),
                ("L2-12", "农户投诉或纠纷怎么处理？",
                 "不清楚。如果要按服务质量扣减分红，就必须有申诉渠道。", "建议答"),
            ]),
            ("三、组织与分工", [
                ("L2-13", "三个端（企业后台、技师端、农户端）分别给谁用、各做什么？",
                 "我们按此设想：后台做管理与审核；技师端接单、打卡、开方；"
                 "农户端下单、查进度、验收、扫码验真。", "必答"),
                ("L2-14", "区域怎么划分？（比如县 → 乡镇 / 网格？）派单按什么范围匹配？",
                 "我们按「县 → 网格」两级设想。", "必答"),
            ]),
        ],
    },
    {
        "name": "L3-规则与算法",
        "title": "第三层：规则与算法",
        "note": "建议 L1、L2 确认后再填。本层确定资质怎么审、分红怎么算、合规怎么落。共 16 题。",
        "groups": [
            ("一、资质与合规", [
                ("L3-01", "技师的资质由谁审、多长时间审一次？",
                 "我们按「企业运营管理员人工审核」设想。", "必答"),
                ("L3-02", "审核时具体看哪些项？",
                 "我们按「证件真伪、有效期、许可经营范围与作业类型是否匹配、发证机关」设想。", "必答"),
                ("L3-03", "许可证到期怎么办？过期后还能接单吗？",
                 "我们按「过期即不能接单，且当月不计分红」设想。", "建议答"),
                ("L3-04", "是否已经具备对接国家 / 省级监管系统的资质和接口？",
                 "目前的说法是「正在申请中」。", "必答"),
                ("L3-05", "一物一码溯源，监管上具体有什么要求？（怎么赋码、要不要上报）",
                 "只知道「市场监管合规上货」，具体要求不清楚。", "建议答"),
                ("L3-06", "需不需要做窜货预警、批次冻结？",
                 "我们目前按「本期不做、二期再说」设想。", "建议答"),
            ]),
            ("二、分红与结算", [
                ("L3-07", "分红的钱从哪里来？",
                 "我们按「只来自服务收入分成」设想——不含卖产品的提成，"
                 "不含团队或公司的利润分配。", "必答"),
                ("L3-08", "现在有没有成文的分成 / 激励制度文件？能不能提供一份？",
                 "之前的说法是「有想法，但还没写成文件」。", "必答"),
                ("L3-09", "「让新农人成立股东」——是真股权（要办工商登记），"
                          "还是只分钱的虚拟分红权？",
                 "我们按「虚拟分红权、不涉及工商登记」设想。这一条差别很大："
                 "真股权要走公司法与工商变更。", "必答"),
                ("L3-10", "分红基数里的「服务费」指什么？包含农资销售金额吗？",
                 "我们按「只算服务部分、不含农资」设想。", "必答"),
                ("L3-11", "农资成本怎么分摊到单张工单上？",
                 "不清楚。这一条不定下来，上面的「只算服务部分」就算不出来。", "必答"),
                ("L3-12", "分红分几档？按什么分档？各档系数是多少？",
                 "我们建议：按滚动 12 个月累计服务费分 4 档，系数 1.00 / 1.10 / 1.20 / 1.30。"
                 "—— 这是我们的建议值，不是既有制度，请务必修正或确认。", "必答"),
                ("L3-13", "结算周期和结算日？（比如月结，次月几号）",
                 "我们按「月结」设想，具体日子没定。", "建议答"),
                ("L3-14", "分红按什么税目代扣代缴个税？",
                 "我们按「劳务报酬所得、由公司代扣代缴」设想。"
                 "这一条建议请财务确认。", "必答"),
                ("L3-15", "有没有单人月度封顶？服务质量差要不要扣减分红？怎么申诉？",
                 "不清楚。", "建议答"),
                ("L3-16", "分红通过什么渠道发放？",
                 "我们按「银行银企直联批量代发」设想。", "建议答"),
            ]),
        ],
    },
    {
        "name": "L4-数据与集成",
        "title": "第四层：数据与集成",
        "note": "最后一层，偏技术细节，也可以请 IT 或对接人填。共 11 题。",
        "groups": [
            ("一、基础数据", [
                ("L4-01", "有没有现成的区域 / 网格划分表？能提供吗？", "不清楚。", "必答"),
                ("L4-02", "技师名录和证件资料现在在哪里？以什么形式提供？", "不清楚。", "必答"),
                ("L4-03", "「30 多个产品」的清单和农药登记证号，能提供吗？", "不清楚。", "建议答"),
                ("L4-04", "历史服务记录需要导入系统吗？", "不清楚。", "建议答"),
                ("L4-05", "各类编号规则（技师编号、工单号、批次号、溯源码）现在怎么编的？",
                 "原型里的编号格式是我们编造的，没有依据。", "建议答"),
            ]),
            ("二、端与集成", [
                ("L4-06", "技师端和农户端用微信小程序吗？账号怎么登录（手机号 / 微信授权）？",
                 "我们按「微信小程序」设想，登录方式未定。", "建议答"),
                ("L4-07", "消息通知走什么渠道？（短信 / 微信模板消息 / App 推送）",
                 "不清楚。", "建议答"),
                ("L4-08", "需要和现有系统对接吗？（ERP、进销存、企业微信等）", "不清楚。", "建议答"),
                ("L4-09", "部署在哪里？（自有服务器 / 云）", "不清楚。", "建议答"),
                ("L4-10", "数据需要保存多久？农户个人信息（手机号、身份证、地块）"
                          "有什么合规要求？",
                 "不清楚。", "建议答"),
                ("L4-11", "首批上线覆盖哪个县、多少名技师、多少农户？",
                 "我们按「单县试点、30–50 名技师」设想。", "必答"),
            ]),
        ],
    },
]


def style_header(ws, row, height=26):
    for idx, (label, width) in enumerate(COLS, start=1):
        cell = ws.cell(row=row, column=idx, value=label)
        cell.font = Font(bold=True, color="FFFFFF", size=10.5)
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[row].height = height


def build_question_sheet(wb, spec):
    ws = wb.create_sheet(spec["name"])
    last_col = get_column_letter(len(COLS))

    # 标题
    ws.merge_cells(f"A1:{last_col}1")
    t = ws["A1"]
    t.value = spec["title"]
    t.font = Font(bold=True, size=14, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=BRAND)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    # 说明
    ws.merge_cells(f"A2:{last_col}2")
    n = ws["A2"]
    n.value = spec["note"]
    n.font = Font(size=10, color="404040")
    n.fill = PatternFill("solid", fgColor=GREY_TINT)
    n.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[2].height = 34

    style_header(ws, 3)

    row = 4
    blocks = []
    for group_name, items in spec["groups"]:
        # 分组标题行
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLS))
        g = ws.cell(row=row, column=1, value=group_name)
        g.font = Font(bold=True, size=11, color=BRAND)
        g.fill = PatternFill("solid", fgColor=BRAND_TINT)
        g.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(1, len(COLS) + 1):
            ws.cell(row=row, column=c).border = BORDER
        ws.row_dimensions[row].height = 22
        row += 1

        block_first = row
        for no, question, ours, prio in items:
            ws.cell(row=row, column=1, value=no).alignment = Alignment(
                horizontal="center", vertical="center")
            ws.cell(row=row, column=2, value=question)
            ws.cell(row=row, column=3, value=ours)

            d = ws.cell(row=row, column=4)
            e = ws.cell(row=row, column=5)
            f = ws.cell(row=row, column=6)
            p = ws.cell(row=row, column=7, value=prio)

            for c in (2, 3):
                ws.cell(row=row, column=c).alignment = Alignment(
                    vertical="top", wrap_text=True)
            for c in (4, 5, 6):
                ws.cell(row=row, column=c).alignment = Alignment(
                    vertical="top", wrap_text=True)
            d.fill = PatternFill("solid", fgColor=FILL_IN)
            e.fill = PatternFill("solid", fgColor=FILL_IN)
            f.fill = PatternFill("solid", fgColor=FILL_IN)
            d.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            p.font = Font(bold=True, size=10, color=PRIORITY_COLOR[prio])
            p.alignment = Alignment(horizontal="center", vertical="center")

            for c in range(1, len(COLS) + 1):
                ws.cell(row=row, column=c).border = BORDER

            # 行高：按「问题」与「我们的理解」的换行数估算
            lines = max(
                ceil(len(question) / 19),
                ceil(len(ours) / 25),
                2,
            )
            ws.row_dimensions[row].height = lines * 15 + 8

            row += 1
        blocks.append((block_first, row - 1))

    # 下拉：是否一致（只加在真实题目行上，不含分组标题的合并单元格）
    dv = DataValidation(
        type="list",
        formula1='"一致,不一致（见E列）,不确定,不适用"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "请从下拉列表中选择"
    dv.errorTitle = "输入无效"
    ws.add_data_validation(dv)
    for f, l in blocks:
        dv.add(f"D{f}:D{l}")

    ws.freeze_panes = "A4"
    return ws, blocks


def build_guide(wb, sheet_ranges):
    ws = wb.create_sheet("填写说明", 0)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 78

    ws.merge_cells("B2:C2")
    ws["B2"] = "华农智服 · 需求澄清与确认问卷"
    ws["B2"].font = Font(bold=True, size=16, color=BRAND)
    ws.row_dimensions[2].height = 32

    ws.merge_cells("B3:C3")
    ws["B3"] = "用途：把我们对需求的理解摆出来，请您确认或纠正。确认过的部分我们才会进入详细需求文档。"
    ws["B3"].font = Font(size=11, color="404040")
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 30

    rows = [
        ("分层说明", "共 4 层，请按顺序填，不要跳层：", True),
        ("", "L1 需求本质（7 题）→ L2 业务范围与流程（14 题）"
             "→ L3 规则与算法（16 题）→ L4 数据与集成（11 题）", False),
        ("", "★ 请先只填 L1。L1 确认后我们再一起看 L2，避免一次面对几十个细节问题。", False),
        ("填写方法", "每题只需两步：", True),
        ("", "第 1 步：看「我们的理解」这一列，在下拉里选「一致」或「不一致」。", False),
        ("", "第 2 步：选「不一致」的，在「请填写你的答案 / 修正」列写出正确做法。", False),
        ("", "「我们的理解」仅供参考，写错的请直接改，不用顾虑。", False),
        ("重要度", "必答 = 不确认无法继续设计；建议答 = 影响方案质量；可选 = 有则更好。", True),
        ("", "如果某一题暂时答不上来，可在「补充说明」里写「待内部确认」，不必空着。", False),
        ("下拉含义", "一致 = 我们的理解和实际一致；不一致 = 请在第 5 列写正确做法；", True),
        ("", "不确定 = 需要再核实；不适用 = 本业务不涉及这一项。", False),
        ("回传方式", "填完后请回传本文件；如有疑问可先在微信里说明，我们改完再发。", True),
        ("", "填表人：____________________    日期：____________________", False),
        ("", "回传截止：____________________", False),
    ]

    r = 5
    for label, text, bold in rows:
        if label:
            c = ws.cell(row=r, column=2, value=label)
            c.font = Font(bold=True, size=11, color=BRAND)
            c.alignment = Alignment(vertical="top")
            c.fill = PatternFill("solid", fgColor=BRAND_TINT)
            c.border = BORDER
        else:
            ws.cell(row=r, column=2).border = BORDER
        v = ws.cell(row=r, column=3, value=text)
        v.font = Font(bold=bold, size=10.5)
        v.alignment = Alignment(wrap_text=True, vertical="center")
        v.border = BORDER
        ws.row_dimensions[r].height = 24 if len(text) < 50 else 34
        r += 1

    # 进度入口提示
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(row=r, column=2,
            value="填写进度可在「填写进度」工作表查看（自动统计，不用手填）。")
    ws.cell(row=r, column=2).font = Font(italic=True, size=10, color="808080")
    return ws


def build_progress(wb, sheet_ranges):
    ws = wb.create_sheet("填写进度")
    headers = ["层次", "题数", "已回复（是否一致已选）", "已写明修正", "完成度"]
    widths = [24, 10, 22, 18, 14]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10.5)
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30

    r = 2
    for label, sheet_name, blocks in sheet_ranges:
        total = sum(l - f + 1 for f, l in blocks)
        cnt_d = "+".join(f"COUNTIF('{sheet_name}'!$D${f}:$D${l},\"<>\")" for f, l in blocks)
        cnt_e = "+".join(f"COUNTIF('{sheet_name}'!$E${f}:$E${l},\"<>\")" for f, l in blocks)
        ws.cell(row=r, column=1, value=label).border = BORDER
        ws.cell(row=r, column=2, value=total).border = BORDER
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=f"={cnt_d}").border = BORDER
        ws.cell(row=r, column=4, value=f"={cnt_e}").border = BORDER
        ws.cell(row=r, column=5,
                value=f"=IF(B{r}=0,0,C{r}/B{r})").border = BORDER
        ws.cell(row=r, column=5).number_format = "0%"
        for c in range(1, 6):
            ws.cell(row=r, column=c).alignment = Alignment(
                horizontal="center", vertical="center")
        r += 1

    ws.cell(row=r, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B2:B{r-1})").font = Font(bold=True)
    ws.cell(row=r, column=3, value=f"=SUM(C2:C{r-1})").font = Font(bold=True)
    ws.cell(row=r, column=4, value=f"=SUM(D2:D{r-1})").font = Font(bold=True)
    ws.cell(row=r, column=5, value=f"=IF(B{r}=0,0,C{r}/B{r})").font = Font(bold=True)
    ws.cell(row=r, column=5).number_format = "0%"
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = BORDER
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    sheet_ranges = []
    for spec in SHEETS:
        _, blocks = build_question_sheet(wb, spec)
        sheet_ranges.append((spec["title"], spec["name"], blocks))

    build_guide(wb, sheet_ranges)
    build_progress(wb, sheet_ranges)

    wb.active = 0
    out = Path(__file__).resolve().parent.parent / "docs" / "华农智服-需求澄清与确认问卷.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    total = sum(l - f + 1 for _, _, blocks in sheet_ranges for f, l in blocks)
    print(f"saved: {out}")
    print(f"sheets: {wb.sheetnames}")
    print(f"questions: {total}")
    for label, name, blocks in sheet_ranges:
        n = sum(l - f + 1 for f, l in blocks)
        print(f"  {label}: {n} (blocks={blocks})")


if __name__ == "__main__":
    main()
